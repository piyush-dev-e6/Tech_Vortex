#!/usr/bin/env python3
"""Deterministic resume text extraction and evidence-backed skill parsing."""

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from skill_normalizer import ALIASES, display_skill, normalize_term

MAX_RESUME_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
SKILL_CATEGORIES = {
    "programming_languages": {"python", "javascript", "typescript", "java", "c", "c++", "c#", "go", "ruby", "php", "rust", "kotlin", "swift", "r", "bash"},
    "frameworks": {"react", "next.js", "vue", "angular", "express", "fastapi", "pytorch", "tensorflow", "scikit-learn", "langchain"},
    "tools": {"docker", "kubernetes", "aws", "azure", "gcp", "git", "terraform", "tableau", "power bi", "pandas", "numpy", "mongodb", "postgresql", "mysql", "redis", "linux", "wireshark", "jenkins", "prometheus", "grafana", "snowflake", "bigquery"},
}
ONTOLOGY_PATH = Path(__file__).with_name("skill_ontology.json")
with ONTOLOGY_PATH.open("r", encoding="utf-8") as source:
    SKILL_ONTOLOGY = json.load(source)
SKILL_VOCABULARY = {normalize_term(skill) for skill in SKILL_ONTOLOGY.get("skills", [])}
ALIAS_TERMS = {re.sub(r"\s+", " ", str(alias).strip().lower()): normalize_term(canonical) for alias, canonical in ALIASES.items()}
MAX_EXTRACTED_TEXT_CHARS = 250_000


class ResumeParseError(ValueError):
    pass


def _extract_text(data: bytes, extension: str) -> str:
    if extension == ".txt":
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ResumeParseError("TXT must be valid UTF-8 text.") from exc
    if extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ResumeParseError("The file does not contain a valid PDF header.")
        try:
            from pypdf import PdfReader
        except ImportError:
            pdftotext = shutil.which("pdftotext")
            if not pdftotext:
                raise ResumeParseError("PDF parsing is unavailable: install pypdf with backend/requirements.txt or install pdftotext.")
            try:
                completed = subprocess.run([pdftotext, "-layout", "-", "-"], input=data, stdout=subprocess.PIPE,
                                           stderr=subprocess.PIPE, timeout=25, check=True)
                return completed.stdout.decode("utf-8", errors="replace")
            except subprocess.TimeoutExpired as exc:
                raise ResumeParseError("PDF text extraction timed out.") from exc
            except (OSError, subprocess.CalledProcessError) as exc:
                raise ResumeParseError("Could not read this PDF. It may be malformed, encrypted, or image-only.") from exc
        try:
            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ResumeParseError("Password-protected PDFs are not supported.")
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ResumeParseError:
            raise
        except Exception as exc:
            raise ResumeParseError("Could not read this PDF. It may be malformed or encrypted.") from exc
    if extension == ".docx":
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ResumeParseError("The DOCX document is missing its main document content.")
                document_xml = archive.read("word/document.xml")
            try:
                from docx import Document
            except ImportError:
                # OOXML is a ZIP/XML format; this fallback keeps basic paragraph
                # and table text available when python-docx is not installed.
                root = ET.fromstring(document_xml)
                namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paragraphs = ["".join(node.text or "" for node in paragraph.findall(".//w:t", namespace))
                              for paragraph in root.findall(".//w:p", namespace)]
                return "\n".join(paragraph for paragraph in paragraphs if paragraph.strip())
            document = Document(BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return "\n".join(paragraphs)
        except ResumeParseError:
            raise
        except Exception as exc:
            raise ResumeParseError("Could not read this DOCX. The file may be malformed.") from exc
    raise ResumeParseError("Only PDF, DOCX, and TXT resumes are supported.")


def _lines(text: str) -> List[str]:
    return [re.sub(r"\s+", " ", line).strip(" \t•*-\u2022") for line in text.replace("\r", "\n").split("\n") if line.strip()]


def _field(value: Optional[Any], confidence: float = 0.0) -> Dict[str, Any]:
    return {"value": value, "confidence": round(confidence, 2), "uncertain": value is None or confidence < 0.75}


SECTION_HEADING_ALIASES = {
    "summary": {"summary", "professional summary", "profile", "professional profile", "objective", "career objective", "about me"},
    "skills": {"skills", "technical skills", "key skills", "core competencies", "technical expertise", "technical toolkit", "toolkit", "tools and technologies", "skills and technologies"},
    "education": {"education", "academic background", "academic qualifications", "educational qualifications", "coursework", "relevant coursework"},
    "experience": {"experience", "work experience", "professional experience", "employment", "internships", "internship", "research experience", "work history", "experience / research exposure", "experience/research exposure"},
    "projects": {"projects", "academic projects", "personal projects", "selected projects", "project experience"},
    "certifications": {"certifications", "certificates", "licenses and certifications", "licenses"},
    "achievements": {"achievements", "awards", "honors and awards", "honours and awards", "accomplishments", "achievements & activities", "achievements and activities"},
    "interests": {"interests", "areas of interest", "professional interests", "extracurricular activities", "activities"},
}
HEADING_TO_SECTION = {alias: section for section, aliases in SECTION_HEADING_ALIASES.items() for alias in aliases}


def _heading_key(value: str) -> str:
    value = re.sub(r"^\s*\d{1,2}[.)]\s*", "", value)
    return re.sub(r"\s+", " ", value.strip().strip(":-–— ")).lower()


def _detect_sections(text: str) -> Tuple[Dict[str, Dict[str, Any]], List[Dict[str, Any]]]:
    sections: Dict[str, Dict[str, Any]] = {}
    entries: List[Dict[str, Any]] = []
    active = "unknown"
    counters: Dict[str, int] = {}
    normalized_text = text.replace("\r\n", "\n").replace("\r", "\n")
    for line_number, raw in enumerate(normalized_text.split("\n"), start=1):
        line = re.sub(r"\s+", " ", raw).strip(" \t•*\u2022-")
        if not line:
            continue
        inline = re.match(r"^([^:]{2,45}):\s*(.+)$", line)
        if inline and _heading_key(inline.group(1)) in HEADING_TO_SECTION:
            active = HEADING_TO_SECTION[_heading_key(inline.group(1))]
            content = inline.group(2).strip()
            if not content:
                continue
            section_line = counters.get(active, 0) + 1
            counters[active] = section_line
            sections.setdefault(active, {"heading": inline.group(1).strip(), "text": "", "lineNumbers": []})
            sections[active]["text"] = "\n".join(part for part in [sections[active]["text"], content] if part)
            sections[active]["lineNumbers"].append(line_number)
            entries.append({"section": active, "text": content, "line": line_number, "sectionLine": section_line})
            continue
        key = _heading_key(line)
        if key in HEADING_TO_SECTION:
            active = HEADING_TO_SECTION[key]
            sections.setdefault(active, {"heading": line.strip().rstrip(":"), "text": "", "lineNumbers": []})
            continue
        section_line = counters.get(active, 0) + 1
        counters[active] = section_line
        sections.setdefault(active, {"heading": "Unclassified text" if active == "unknown" else active.title(), "text": "", "lineNumbers": []})
        sections[active]["text"] = "\n".join(part for part in [sections[active]["text"], line] if part)
        sections[active]["lineNumbers"].append(line_number)
        entries.append({"section": active, "text": line, "line": line_number, "sectionLine": section_line})
    return sections, entries


def _section_lines(lines: List[str], headings: Tuple[str, ...]) -> List[str]:
    sections, _ = _detect_sections("\n".join(lines))
    wanted = {HEADING_TO_SECTION.get(_heading_key(heading), heading) for heading in headings}
    return [line for section, data in sections.items() if section in wanted for line in data["text"].splitlines()]

def _extract_name(lines: List[str]) -> Tuple[Optional[str], float]:
    excluded = re.compile(r"resume|curriculum vitae|email|phone|linkedin|github|engineer|developer|candidate", re.I)
    for line in lines[:4]:
        candidate = line.split("|")[0].strip()
        candidate = re.sub(r"\s+(?:[-–|].*)$", "", candidate).strip()
        words = candidate.split()
        if 2 <= len(words) <= 4 and not excluded.search(candidate) and not re.search(r"[@\d]", candidate):
            if all(re.fullmatch(r"[A-Za-z][A-Za-z.'’-]*", word) for word in words):
                return candidate.title(), 0.90 if line == candidate else 0.84
    return None, 0.0


def _extract_resume_fields(text: str) -> Dict[str, Any]:
    lines = _lines(text)
    name, name_confidence = _extract_name(lines)
    whole = "\n".join(lines)
    sections, entries = _detect_sections(text)

    degree_patterns = [
        r"\b(B\.?\s*Tech\.?|B\.?\s*E\.?|B\.?\s*Sc\.?|B\.?\s*CA|M\.?\s*Tech\.?|M\.?\s*E\.?|M\.?\s*Sc\.?|M\.?\s*CA|Ph\.?\s*D\.?)\b(?:\s+(?:in|of)\s+[^,;|\n]+)?",
        r"\b(?:Bachelor|Master|Doctorate|Associate)(?:'s)?\s+(?:of|in)\s+[^,;|\n]+",
    ]
    degree = None
    for line in lines:
        for pattern in degree_patterns:
            match = re.search(pattern, line, re.I)
            if match:
                degree = re.sub(r"\s+", " ", match.group(0)).strip(" ,.;")
                break
        if degree:
            break
    degree_confidence = 0.90 if degree else 0.0

    institution = None
    for line in lines:
        match = re.search(r"\b(?:[A-Z][\w.&'-]*\s+){0,6}(?:University|College|Institute of Technology|Institute|IIT|NIT)\b[^,;|]*", line, re.I)
        if match:
            institution = match.group(0).strip(" ,.;")
            break
    institution_confidence = 0.82 if institution else 0.0

    cgpa_match = re.search(r"\b(?:CGPA|GPA|CPI)\s*[:=\-]?\s*(\d{1,2}(?:\.\d{1,2})?)\s*(?:/\s*(?:10|4))?", whole, re.I)
    cgpa = float(cgpa_match.group(1)) if cgpa_match else None
    cgpa_confidence = 0.95 if cgpa is not None and cgpa <= 10 else 0.0
    if cgpa is not None and cgpa > 10:
        cgpa = None

    graduation_year = None
    year_patterns = [r"(?:graduat(?:ion|ed)|expected\s+(?:to\s+)?graduate|passing\s+year|batch)\D{0,24}(20\d{2})", r"\b(20\d{2})\s+(?:batch|graduate|graduation)\b"]
    for pattern in year_patterns:
        match = re.search(pattern, whole, re.I)
        if match:
            graduation_year = int(match.group(1))
            break
    year_confidence = 0.86 if graduation_year else 0.0

    section_quality = {"projects": 0.78, "experience": 0.80, "skills": 0.68, "certifications": 0.72, "achievements": 0.70, "education": 0.60, "summary": 0.60, "interests": 0.48, "unknown": 0.48}
    action_pattern = re.compile(r"\b(?:built|developed|implemented|designed|trained|optimized|deployed|created|evaluated|analyzed|maintained|integrated|engineered|automated|configured|led|authored|applied)\b", re.I)
    match_terms = set(SKILL_VOCABULARY) | set(ALIAS_TERMS)
    skill_evidence: Dict[str, Dict[str, Any]] = {}
    for entry in entries:
        line = entry["text"]
        found_here: Dict[str, List[str]] = {}
        for term in sorted(match_terms, key=lambda value: (-len(value), value)):
            canonical = normalize_term(ALIAS_TERMS.get(term, term))
            if not canonical:
                continue
            if len(term) <= 2 and term not in {"c#", "c+"} and entry["section"] != "skills":
                continue
            match = re.search(r"(?<![\w])" + re.escape(term) + r"(?![\w])", line, re.I)
            if match:
                found_here.setdefault(canonical, []).append(match.group(0))
        for canonical, aliases in found_here.items():
            record = skill_evidence.setdefault(canonical, {"evidence": [], "aliasesDetected": []})
            evidence_key = (entry["section"], line)
            if not any((item["section"], item["text"]) == evidence_key for item in record["evidence"]):
                record["evidence"].append({
                    "section": entry["section"],
                    "text": line[:500],
                    "location": {"line": entry["line"], "sectionLine": entry["sectionLine"]},
                })
            for alias in aliases:
                if alias not in record["aliasesDetected"]:
                    record["aliasesDetected"].append(alias)

    skills = []
    for canonical, record in sorted(skill_evidence.items()):
        evidence = record["evidence"]
        scores = []
        for item in evidence:
            score = section_quality.get(item["section"], 0.48)
            if item["section"] in {"projects", "experience"} and action_pattern.search(item["text"]):
                score += 0.12
            scores.append(score)
        repeated_bonus = min(0.08, max(0, len(evidence) - 1) * 0.025)
        section_bonus = 0.04 if len({item["section"] for item in evidence}) > 1 else 0.0
        confidence = round(min(0.99, max(scores, default=0.0) + repeated_bonus + section_bonus), 2)
        category = next((category for category, values in SKILL_CATEGORIES.items() if canonical in values), "technical_skills")
        label = "strong" if confidence >= 0.82 else "moderate" if confidence >= 0.65 else "limited"
        strongest = max(evidence, key=lambda item: (section_quality.get(item["section"], 0.48) + (0.12 if item["section"] in {"projects", "experience"} and action_pattern.search(item["text"]) else 0), -item["location"]["line"]))
        skills.append({
            "canonical": canonical,
            "canonicalName": display_skill(canonical),
            "name": display_skill(canonical),
            "display": display_skill(canonical),
            "confidence": confidence,
            "confidenceLabel": label,
            "confidenceMethod": "Deterministic evidence-strength heuristic; not a calibrated probability.",
            "source": "resume",
            "confirmed": False,
            "aliasesDetected": sorted(set(record["aliasesDetected"]), key=lambda value: (value.lower(), value)),
            "evidence": evidence,
            "evidenceText": strongest["text"],
            "category": category,
            "uncertain": confidence < 0.65,
        })

    projects = [entry["text"] for entry in entries if entry["section"] == "projects"]
    experience = [entry["text"] for entry in entries if entry["section"] == "experience"]
    certifications = [entry["text"] for entry in entries if entry["section"] == "certifications"]
    education_value = {"degree": degree, "institution": institution, "graduationYear": graduation_year, "cgpa": cgpa}
    education_confidence = max(degree_confidence, institution_confidence, year_confidence, cgpa_confidence)
    field_confidence = {
        "name": _field(name, name_confidence), "education": _field(education_value if education_confidence else None, education_confidence),
        "degree": _field(degree, degree_confidence), "college": _field(institution, institution_confidence),
        "graduationYear": _field(graduation_year, year_confidence), "cgpa": _field(cgpa, cgpa_confidence),
    }
    return {
        "name": name,
        "education": education_value if education_confidence else None,
        "degree": degree,
        "college": institution,
        "graduationYear": graduation_year,
        "cgpa": cgpa,
        "skills": skills,
        "programmingLanguages": [skill for skill in skills if skill["category"] == "programming_languages"],
        "frameworks": [skill for skill in skills if skill["category"] == "frameworks"],
        "tools": [skill for skill in skills if skill["category"] == "tools"],
        "technicalSkills": [skill for skill in skills if skill["category"] == "technical_skills"],
        "projects": projects,
        "experience": experience,
        "internshipsOrWorkExperience": experience,
        "certifications": certifications,
        "sections": sections,
        "fieldConfidence": field_confidence,
        "uncertainFields": [key for key, value in field_confidence.items() if value["uncertain"]],
        "confidenceMethod": "Deterministic rule-based field and skill evidence; confidence values are heuristic, not calibrated probabilities.",
        "warnings": [] if text.strip() else ["No readable text was found in the resume; no fields or skills were inferred."],
    }

def parse_resume_bytes(data: bytes, filename: str) -> Dict[str, Any]:
    extension = Path(filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeParseError("Resume exceeds the 5 MB upload limit.")
    if extension == ".docx" and not data.startswith(b"PK"):
        raise ResumeParseError("The file does not appear to be a valid DOCX document.")
    text = _extract_text(data, extension)
    candidate_profile = _extract_resume_fields(text)
    content_hash = hashlib.sha256(data).hexdigest()
    return {
        "success": True,
        "filename": Path(filename).name,
        "fileType": extension[1:],
        "extractedTextLength": len(text),
        "textPreview": text.strip()[:3000],
        "extractedText": text[:MAX_EXTRACTED_TEXT_CHARS],
        "extractedTextTruncated": len(text) > MAX_EXTRACTED_TEXT_CHARS,
        "contentHash": content_hash,
        "resume": {"fileName": Path(filename).name, "status": "analyzed", "contentHash": content_hash},
        "sections": candidate_profile["sections"],
        "candidates": candidate_profile["skills"],
        "candidateProfile": candidate_profile,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract structured resume information")
    parser.add_argument("--file", required=True)
    parser.add_argument("--filename", required=True)
    args = parser.parse_args()
    try:
        data = Path(args.file).read_bytes()
        result = parse_resume_bytes(data, args.filename)
        print(json.dumps(result, ensure_ascii=False))
    except ResumeParseError as error:
        print(json.dumps({"success": False, "error": str(error)}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"success": False, "error": "Resume parsing failed safely."}, ensure_ascii=False), file=sys.stdout)


if __name__ == "__main__":
    main()
